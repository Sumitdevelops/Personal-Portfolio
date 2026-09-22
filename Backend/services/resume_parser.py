from pathlib import Path
from docx import Document


def resume_extract_text(path: Path):
    doc = Document(path)

    text = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                row_text.append(cell.text)

            text.append(" | ".join(row_text))

    return "\n".join(text)